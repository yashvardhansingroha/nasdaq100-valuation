"""Generate Word summary of the Meta-Portfolio case study."""

from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor

OUT = Path(__file__).resolve().parent / "reports" / "Multi_Strategy_Portfolio_Allocation_Case_Study.docx"


def set_run_font(run, size=11, bold=False, color=None):
    run.font.name = "Calibri"
    run._element.rPr.rFonts.set(qn("w:eastAsia"), "Calibri")
    run.font.size = Pt(size)
    run.bold = bold
    if color:
        run.font.color.rgb = RGBColor(*color)


def main() -> None:
    doc = Document()
    for section in doc.sections:
        section.top_margin = Inches(1)
        section.bottom_margin = Inches(1)
        section.left_margin = Inches(1)
        section.right_margin = Inches(1)

    def add_heading_custom(text, level=1):
        p = doc.add_paragraph()
        run = p.add_run(text)
        if level == 0:
            set_run_font(run, 18, True, (20, 40, 80))
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        elif level == 1:
            set_run_font(run, 14, True, (20, 40, 80))
            p.space_before = Pt(14)
            p.space_after = Pt(6)
        else:
            set_run_font(run, 12, True, (20, 40, 80))
            p.space_before = Pt(10)
            p.space_after = Pt(4)
        return p

    def add_body(text):
        p = doc.add_paragraph()
        run = p.add_run(text)
        set_run_font(run, 11, False)
        p.paragraph_format.space_after = Pt(8)
        p.paragraph_format.line_spacing = 1.15
        return p

    def add_bullet(text):
        p = doc.add_paragraph(style="List Bullet")
        run = p.add_run(text)
        set_run_font(run, 11, False)
        p.paragraph_format.space_after = Pt(3)
        return p

    def add_table(headers, rows):
        table = doc.add_table(rows=1 + len(rows), cols=len(headers))
        table.style = "Table Grid"
        for i, h in enumerate(headers):
            cell = table.rows[0].cells[i]
            cell.text = ""
            run = cell.paragraphs[0].add_run(h)
            set_run_font(run, 10, True)
        for r_i, row in enumerate(rows):
            for c_i, val in enumerate(row):
                cell = table.rows[r_i + 1].cells[c_i]
                cell.text = ""
                run = cell.paragraphs[0].add_run(str(val))
                set_run_font(run, 10, False)
        doc.add_paragraph()

    def center_line(text, size=11, bold=False, color=(80, 80, 80)):
        p = doc.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        run = p.add_run(text)
        set_run_font(run, size, bold, color)
        return p

    add_heading_custom("Multi-Strategy Portfolio Allocation Case Study", 0)
    center_line("Quant Club, BITS Pilani  |  Portfolio Construction & Risk")
    center_line("Author: Yashvardhan  |  August 2026")
    center_line("Sample period: January 2019 – August 2026  (~1,907 trading days)", 10)

    add_heading_custom("1. Project Overview", 1)
    add_body(
        "This project studies how capital should be allocated across distinct trading-strategy "
        "sleeves rather than across individual stocks. It models a multi-strategy allocator that "
        "dynamically redistributes funds based on recent risk, then measures whether that approach "
        "improves outcomes versus simple static benchmarks."
    )
    add_body(
        "The central research question is: when market regimes change, how should an allocator "
        "budget risk across momentum, defensive, and crisis exposures — and when does a permanent "
        "volatility sleeve help versus hurt?"
    )

    add_heading_custom("2. What the Project Does", 1)
    add_body("In practical terms, the project:")
    add_bullet(
        "Pulls real market history for three public ETF strategy proxies: MTUM (momentum), "
        "USMV (minimum volatility / defensive), and VIXY (crisis / volatility)."
    )
    add_bullet(
        "Estimates risk on a rolling 60-day window — annualized volatility, covariance, and "
        "correlation — without looking ahead in the data."
    )
    add_bullet(
        "Computes three dynamic allocation methods each day: inverse-volatility risk parity, "
        "minimum-variance optimization, and maximum diversification."
    )
    add_bullet(
        "Applies a realistic 10 basis-point transaction cost on portfolio turnover for all "
        "dynamic methods."
    )
    add_bullet(
        "Benchmarks results against a static equal-weight portfolio (one-third each) and a "
        "static 60/40 portfolio (60% MTUM / 40% USMV / 0% VIXY)."
    )
    add_bullet(
        "Reports performance metrics (return, volatility, Sharpe, Sortino, max drawdown, "
        "Calmar, VaR), attribution by sleeve, and outcomes by market regime "
        "(COVID crash, 2021 bull, 2022 rate shock, 2023 recovery)."
    )

    add_heading_custom("3. Strategy Sleeves", 1)
    add_table(
        ["Sleeve", "ETF", "Economic Role"],
        [
            ["Momentum", "MTUM", "Equity factor tilt; participates in trending markets"],
            ["Defensive", "USMV", "Low-beta / min-vol; capital preservation in risk-off periods"],
            ["Crisis / Vol", "VIXY", "Convexity in market panics; carry cost in calm markets"],
        ],
    )

    add_heading_custom("4. Allocation Methods", 1)
    add_bullet(
        "Inverse Volatility (InvVol): capital weight inversely proportional to recent volatility; "
        "fully invested and long-only."
    )
    add_bullet(
        "Minimum Variance (MinVar): minimizes portfolio variance using the rolling covariance "
        "matrix, with each sleeve constrained between 5% and 60%."
    )
    add_bullet(
        "Maximum Diversification (MaxDiv): maximizes the diversification ratio under the same bounds."
    )
    add_bullet("Equal Weight: permanent one-third allocation to each sleeve (static benchmark).")
    add_bullet(
        "60/40: permanent 60% MTUM / 40% USMV with no VIXY (return-seeking static benchmark)."
    )

    add_heading_custom("5. Key Outcomes (Net of Transaction Costs)", 1)
    add_body(
        "Results below are from the full sample backtest after deducting 10bps of friction "
        "on dynamic rebalancing."
    )
    add_table(
        ["Method", "Ann. Return", "Sharpe", "Max Drawdown", "VaR 95% (1-day)"],
        [
            ["Inverse Volatility", "7.8%", "0.30", "-22.8%", "-0.86%"],
            ["Minimum Variance", "5.0%", "0.05", "-28.0%", "-0.79%"],
            ["Max Diversification", "3.8%", "-0.07", "-36.6%", "-0.87%"],
            ["Equal Weight", "-4.6%", "-0.50", "-57.0%", "-1.46%"],
            ["60/40 (no VIXY)", "15.3%", "0.56", "-33.6%", "-1.70%"],
        ],
    )

    add_heading_custom("Main findings", 2)
    add_bullet(
        "Inverse-volatility allocation improved risk-adjusted performance versus equal-weight: "
        "Sharpe rose from -0.50 to 0.30, and maximum drawdown fell from -57% to about -23%."
    )
    add_bullet(
        "Static 60/40 delivered the highest annualized return (~15%) because it never holds VIXY "
        "and therefore avoids the long-run roll / contango cost of a permanent volatility sleeve."
    )
    add_bullet(
        "Crisis convexity is regime-dependent: VIXY helps in sharp volatility spikes "
        "(for example, March 2020) but is expensive insurance in calm markets and in grinding "
        "rate-driven bears such as 2022."
    )
    add_bullet(
        "The practical conclusion is mandate-dependent: use dynamic inverse-vol when the goal is "
        "risk-aware allocation with a crisis sleeve; use static 60/40 when the goal is return "
        "and permanent hedges are optional."
    )

    add_heading_custom("6. Regime Perspective", 1)
    add_bullet(
        "COVID crash (Feb–Mar 2020): the crisis sleeve exhibited strong convexity; "
        "equal-weight benefited from VIXY while 60/40 sold off without a hedge."
    )
    add_bullet(
        "2021 bull: equity-heavy sleeves performed well; carrying VIXY reduced relative returns."
    )
    add_bullet(
        "2022 rate shock: broad risk-off; dynamic risk cuts helped control drawdowns, but the "
        "vol sleeve was less decisive than in a sharp panic."
    )
    add_bullet(
        "2023 recovery: risk-on returns favoured equity blends without a permanent crisis sleeve."
    )

    add_heading_custom("7. Limitations", 1)
    add_bullet(
        "VIXY is a futures-based product with structural roll costs — not a free permanent hedge."
    )
    add_bullet("ETF proxies approximate strategy styles; they are not live hedge-fund P&L.")
    add_bullet("Daily rebalancing overstates typical institutional turnover.")
    add_bullet(
        "The model is long-only and fully invested, with no leverage, shorting, or cash sleeve."
    )
    add_bullet("Past regime results are not forecasts.")

    add_heading_custom("8. Project Deliverables", 1)
    add_bullet("Investment memo summarizing thesis, methodology, results, and recommendation.")
    add_bullet(
        "Excel workbook with returns, dynamic weights, performance metrics, regimes, and attribution."
    )
    add_bullet("Strategy rationale documenting the economic role of each sleeve.")
    add_bullet(
        "Interactive dashboard and charts comparing equity curves and capital weights over time."
    )

    p = doc.add_paragraph()
    run = p.add_run("Educational / research case study. Not investment advice.")
    set_run_font(run, 9, False, (100, 100, 100))
    p.space_before = Pt(18)

    OUT.parent.mkdir(parents=True, exist_ok=True)
    doc.save(OUT)
    print(OUT)
    print("size", OUT.stat().st_size)


if __name__ == "__main__":
    main()
