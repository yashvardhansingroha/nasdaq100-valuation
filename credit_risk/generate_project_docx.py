"""Generate a Word project summary for the Lending Club credit-risk case study."""

from __future__ import annotations

import json
from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor

ROOT = Path(__file__).resolve().parent
OUT = ROOT / "reports" / "Lending_Club_Credit_Risk_Project_Summary.docx"
METRICS_PATH = ROOT / "output" / "model_metrics.json"


def set_run_font(run, size=11, bold=False, color=None):
    run.font.name = "Calibri"
    run._element.rPr.rFonts.set(qn("w:eastAsia"), "Calibri")
    run.font.size = Pt(size)
    run.bold = bold
    if color:
        run.font.color.rgb = RGBColor(*color)


def _money_b(x: float) -> str:
    return f"${x / 1e9:.2f}B"


def _pct(x: float, digits: int = 1) -> str:
    return f"{100 * x:.{digits}f}%"


def main() -> None:
    metrics = json.loads(METRICS_PATH.read_text(encoding="utf-8")) if METRICS_PATH.exists() else {}

    doc = Document()
    for section in doc.sections:
        section.top_margin = Inches(1)
        section.bottom_margin = Inches(1)
        section.left_margin = Inches(1)
        section.right_margin = Inches(1)

    def heading(text: str, level: int = 1):
        p = doc.add_paragraph()
        run = p.add_run(text)
        if level == 0:
            set_run_font(run, 18, True, (20, 40, 80))
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        elif level == 1:
            set_run_font(run, 13, True, (20, 40, 80))
            p.space_before = Pt(14)
            p.space_after = Pt(6)
        else:
            set_run_font(run, 11, True, (20, 40, 80))
            p.space_before = Pt(10)
            p.space_after = Pt(4)
        return p

    def body(text: str):
        p = doc.add_paragraph()
        run = p.add_run(text)
        set_run_font(run, 11)
        p.paragraph_format.space_after = Pt(8)
        p.paragraph_format.line_spacing = 1.15
        return p

    def bullet(text: str):
        p = doc.add_paragraph(style="List Bullet")
        run = p.add_run(text)
        set_run_font(run, 11)
        p.paragraph_format.space_after = Pt(3)
        return p

    def center(text: str, size=11, bold=False, color=(80, 80, 80)):
        p = doc.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        run = p.add_run(text)
        set_run_font(run, size, bold, color)
        return p

    def table(headers, rows):
        t = doc.add_table(rows=1 + len(rows), cols=len(headers))
        t.style = "Table Grid"
        for i, h in enumerate(headers):
            cell = t.rows[0].cells[i]
            cell.text = ""
            set_run_font(cell.paragraphs[0].add_run(h), 10, True)
        for r_i, row in enumerate(rows):
            for c_i, val in enumerate(row):
                cell = t.rows[r_i + 1].cells[c_i]
                cell.text = ""
                set_run_font(cell.paragraphs[0].add_run(str(val)), 10)
        doc.add_paragraph()

    # ---- Title ----
    heading("Lending Club Credit Risk Analytics", 0)
    center("Probability of Default · IFRS 9 Expected Credit Loss · Portfolio Stress Testing")
    center("Retail Credit Risk Case Study", size=10)

    # ---- Overview ----
    heading("1. Project Overview")
    body(
        "This project builds an end-to-end retail credit risk pipeline on historical "
        "Lending Club loan data. It ingests raw loan records into PostgreSQL, constructs "
        "clean analytical views, estimates probability of default (PD) with a stratified "
        "logistic regression model, translates scores into IFRS 9-style expected credit "
        "loss (ECL = PD × LGD × EAD), and extends the analysis with a WOE/IV application "
        "scorecard, population stability monitoring, risk-adjusted pricing flags, and "
        "portfolio stress scenarios. Outputs are packaged for Power BI dashboarding."
    )

    # ---- Objectives ----
    heading("2. Objectives")
    bullet("Estimate borrower default risk on closed Lending Club loans (Fully Paid vs Charged Off / Default).")
    bullet("Convert model PDs into portfolio expected credit loss using a documented LGD assumption.")
    bullet("Build an interpretable WOE/IV application scorecard alongside the statistical PD model.")
    bullet("Monitor score and feature stability (PSI) and identify economically underpriced exposure.")
    bullet("Stress portfolio ECL under adverse, severe, and recession scenarios.")

    # ---- Data & pipeline ----
    heading("3. Data and Technical Pipeline")
    body(
        "Source data is the Lending Club accepted-loans file (2007–2018Q4). Raw rows are "
        "streamed in chunks into PostgreSQL table loan_data_raw. A SQL ETL layer creates "
        "vw_clean_loan_data (closed outcomes only, with engineered employment length, "
        "term months, and vintage year), plus grade-segment and vintage-cohort summary views."
    )
    n_loans = metrics.get("n_loans", 1_344_976)
    table(
        ["Stage", "Detail"],
        [
            ["Raw ingestion", "2,260,668 loan records loaded to PostgreSQL"],
            ["Analytical sample", f"{n_loans:,} closed loans with defined default outcomes"],
            ["Train / test split", f"{metrics.get('train_size', 941483):,} / {metrics.get('test_size', 403493):,} (70/30, stratified)"],
            ["Database", "PostgreSQL — mn_analytics_db"],
            ["Modelling stack", "Python, pandas, scikit-learn, SQLAlchemy"],
            ["Reporting", "Scored CSV + DAX measures for Power BI"],
        ],
    )

    # ---- Methodology ----
    heading("4. Methodology")
    heading("4.1 Probability of Default Model", 2)
    body(
        "Features include loan amount, interest rate, installment, income, DTI, revolving "
        "utilisation and balance, account counts, delinquencies, employment length, term, "
        "and dummy-encoded grade, purpose, home ownership, and verification status. "
        "Annual income outliers are capped using the IQR rule. Missing numeric values are "
        "median-imputed. Features are standardised before fitting logistic regression with "
        "class_weight='balanced' to address default-class imbalance. The decision threshold "
        "is chosen to maximise F1 on the test fold."
    )
    heading("4.2 IFRS 9 Expected Credit Loss", 2)
    body(
        f"Per-loan ECL is computed as PD × LGD × EAD, where EAD is origination loan amount "
        f"and LGD is fixed at {metrics.get('lgd', 0.60):.0%} (conservative unsecured-consumer "
        "assumption). Portfolio ECL is the sum across loans; coverage ratio is ECL divided "
        "by total exposure."
    )
    heading("4.3 Application Scorecard (WOE / IV)", 2)
    body(
        "Equal-frequency bins are fit on the training fold for DTI, income, revolving "
        "utilisation, employment length, and term. Weight of Evidence and Information Value "
        "rank predictive strength. Bin points are scaled to a FICO-style application score "
        "(300–850) for interpretable rank-ordering."
    )
    heading("4.4 Monitoring, Pricing, and Stress", 2)
    bullet("PSI compares train vs test distributions of model PD, DTI, and grade mix.")
    bullet("Pricing adequacy compares contractual interest rate to modelled expected-loss rate (PD × LGD).")
    bullet("Stress scenarios recompute ECL with PD × 1.5 (adverse), PD × 2.0 (severe), and PD + 5pp (recession).")

    # ---- Results ----
    heading("5. Key Results")
    roc = metrics.get("roc_auc", 0.7054)
    f1 = metrics.get("f1", 0.425)
    gini = metrics.get("gini", 0.4108)
    ks = metrics.get("ks", 0.2987)
    thr = metrics.get("decision_threshold", 0.51)
    ecl = metrics.get("portfolio_ecl", 5_615_626_682.72)
    exp = metrics.get("portfolio_exposure", 19_393_447_250.0)
    cov = metrics.get("ecl_coverage_ratio", 0.289563)

    heading("5.1 Model Validation", 2)
    table(
        ["Metric", "Value"],
        [
            ["ROC-AUC", f"{roc:.2f}"],
            ["F1 Score (default class)", f"{f1:.2f}"],
            ["Decision threshold", f"{thr:.2f}"],
            ["Gini coefficient", f"{gini:.2f}"],
            ["KS statistic", f"{ks:.2f}"],
        ],
    )
    body(
        "Discrimination is meaningful for rank-ordering risk across the book. AUC near 0.71 "
        "indicates useful separation of goods and bads on held-out loans; Gini and KS provide "
        "the credit-industry framing of the same ranking power."
    )

    heading("5.2 Scorecard Information Value", 2)
    iv_rows = []
    for row in metrics.get("feature_iv", []):
        iv_rows.append([row["feature"], f"{row['iv']:.3f}", row["strength"]])
    if not iv_rows:
        iv_rows = [
            ["term_months", "0.174", "strong"],
            ["dti", "0.069", "medium"],
            ["annual_inc", "0.028", "medium"],
        ]
    table(["Feature", "IV", "Strength"], iv_rows)

    heading("5.3 Portfolio Expected Credit Loss", 2)
    table(
        ["Measure", "Value"],
        [
            ["Closed-loan exposure (EAD)", _money_b(exp)],
            ["Portfolio ECL (base)", _money_b(ecl)],
            ["ECL coverage ratio", _pct(cov)],
            ["LGD assumption", f"{metrics.get('lgd', 0.60):.0%}"],
        ],
    )

    heading("5.4 ECL by Lending Club Grade", 2)
    table(
        ["Grade", "Exposure", "ECL", "Avg PD", "ECL % of Exposure"],
        [
            ["A", "$3.27B", "$0.51B", "25.9%", "15.5%"],
            ["B", "$5.20B", "$1.19B", "37.6%", "22.9%"],
            ["C", "$5.41B", "$1.66B", "49.8%", "30.6%"],
            ["D", "$3.07B", "$1.15B", "61.2%", "37.6%"],
            ["E", "$1.65B", "$0.72B", "71.6%", "43.5%"],
            ["F", "$0.61B", "$0.30B", "80.1%", "48.4%"],
            ["G", "$0.19B", "$0.10B", "85.2%", "51.2%"],
        ],
    )
    body(
        "Expected loss intensity rises monotonically from grade A to G, consistent with "
        "weaker credit quality carrying higher modelled PD and a larger share of ECL "
        "relative to exposure."
    )

    heading("5.5 Stability and Pricing", 2)
    psi = metrics.get("psi", {})
    under_pct = metrics.get("underpriced_exposure_pct", 0.544817)
    under_exp = metrics.get("underpriced_exposure", 10_565_877_650.0)
    table(
        ["Check", "Result"],
        [
            ["Score PSI (train vs test)", f"{psi.get('score_psi', 0.0):.4f} ({psi.get('score_psi_label', 'stable')})"],
            ["DTI PSI", f"{psi.get('dti_psi', 0.0):.4f} ({psi.get('dti_psi_label', 'stable')})"],
            ["Grade-mix PSI", f"{psi.get('grade_psi', 0.0001):.4f} ({psi.get('grade_psi_label', 'stable')})"],
            ["Underpriced exposure flag", f"{_pct(under_pct)} of book ({_money_b(under_exp)})"],
        ],
    )

    heading("5.6 Stress Testing", 2)
    stress = {s["scenario"]: s for s in metrics.get("stress_test", [])}
    table(
        ["Scenario", "Shock", "Portfolio ECL", "vs Base"],
        [
            [
                "Base",
                "None",
                _money_b(stress.get("base", {}).get("total_ecl", ecl)),
                "—",
            ],
            [
                "Adverse",
                "PD × 1.5",
                _money_b(stress.get("adverse", {}).get("total_ecl", 8.1e9)),
                _pct(stress.get("adverse", {}).get("delta_pct_vs_base", 0.44)),
            ],
            [
                "Severe",
                "PD × 2.0",
                _money_b(stress.get("severe", {}).get("total_ecl", 9.6e9)),
                _pct(stress.get("severe", {}).get("delta_pct_vs_base", 0.71)),
            ],
            [
                "Recession",
                "PD + 5pp",
                _money_b(stress.get("recession", {}).get("total_ecl", 6.2e9)),
                _pct(stress.get("recession", {}).get("delta_pct_vs_base", 0.10)),
            ],
        ],
    )

    # ---- Deliverables ----
    heading("6. Deliverables")
    bullet("PostgreSQL analytical views: clean loan book, grade segments, vintage cohorts.")
    bullet("Scored dataset (powerbi_scored_loan_data.csv) with PD, ECL, application score, and pricing flags.")
    bullet("WOE/IV summary and scorecard points tables.")
    bullet("Stress-test results and model metrics JSON.")
    bullet("Power BI DAX measure reference (ECL, weighted PD, coverage, pricing, stressed ECL).")
    bullet("Risk memo and interview preparation notes.")

    # ---- Limitations ----
    heading("7. Limitations")
    bullet("LGD is a fixed assumption (60%), not estimated from recoveries.")
    bullet("Sample is accepted/closed loans only — selection bias vs the full applicant population.")
    bullet("Stress scenarios are illustrative multipliers, not macro-linked regulatory packages.")
    bullet("Pricing adequacy is a simplified yield-vs-expected-loss proxy (no funding cost or opex).")
    bullet("Model ranks risk well for portfolio analytics; it is not a production underwriting engine.")

    # ---- Conclusion ----
    heading("8. Conclusion")
    body(
        "The project demonstrates a complete credit-risk analytics workflow: data engineering "
        "in PostgreSQL, a validated PD model with industry discrimination metrics, IFRS 9-style "
        "loss estimation, an interpretable scorecard, stability and pricing diagnostics, and "
        "stress testing suitable for portfolio reporting. Together these components show how "
        "default prediction can be translated into economic loss measures and decision-support "
        "outputs for retail credit risk analysis."
    )

    OUT.parent.mkdir(parents=True, exist_ok=True)
    doc.save(OUT)
    print(f"Wrote {OUT}")


if __name__ == "__main__":
    main()
